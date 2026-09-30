"""Validação da versão v2 no navegador (local ou publicada) + capturas de tela + OS de exemplo.
Uso: python tests/capturas_v2.py URL_V2 PREFIXO   (ex.: http://localhost:4174/v2/ local_v2)"""
import sys
from pathlib import Path

import openpyxl
from playwright.sync_api import sync_playwright

url = sys.argv[1].rstrip('/') + '/'
pref = sys.argv[2] if len(sys.argv) > 2 else 'local_v2'
out = Path(__file__).resolve().parents[1] / 'docs' / 'screenshots'
out.mkdir(parents=True, exist_ok=True)
erros, falhas = [], []
VIEWPORTS = [(1920, 1080), (1366, 768), (1024, 768), (390, 844)]


def ok(cond, msg):
    print(('OK   ' if cond else 'FALHA') + ' ' + msg)
    if not cond:
        falhas.append(msg)


def num(t):
    return int(str(t).strip().split()[0].replace('.', '') or 0)


def sem_rolagem(pg, onde):
    r = pg.evaluate("""() => ({h: document.documentElement.scrollHeight, H: innerHeight, w: document.documentElement.scrollWidth, W: innerWidth,
                              bh: document.body.scrollHeight})""")
    ok(r['h'] <= r['H'] and r['w'] <= r['W'] and r['bh'] <= r['H'], f"{onde}: sem rolagem da página ({r['w']}x{r['h']} em {r['W']}x{r['H']})")


def total(pg, col):
    """valor da linha Total na coluna cujo cabeçalho começa com col"""
    return pg.evaluate("""(col) => { const ths = [...document.querySelectorAll('#v-tabela thead th')];
        const i = ths.findIndex((t) => t.childNodes[0].textContent.trim() === col); if (i < 0) return null;
        return document.querySelectorAll('#v-tabela tfoot td')[i].textContent; }""", col)


with sync_playwright() as p:
    b = p.chromium.launch(executable_path='/usr/bin/google-chrome', headless=True)
    ctx = b.new_context(viewport={'width': 1440, 'height': 900}, locale='pt-BR', accept_downloads=True)
    pg = ctx.new_page()
    pg.on('console', lambda m: erros.append(m.text) if m.type == 'error' else None)
    pg.on('pageerror', lambda e: erros.append(str(e)))
    pg.goto(url, wait_until='networkidle')
    pg.wait_for_selector('#v-tabela .gt')
    frota = pg.evaluate("fetch('../data/frota.json').then(r => r.json())")
    meta = pg.evaluate("fetch('../data/meta.json').then(r => r.json())")
    nveic = len(frota['veiculos'])

    # ---------------- Aba 1: Visão geral ----------------
    cards = {c.get_attribute('data-card'): num(c.query_selector('.val').inner_text()) for c in pg.query_selector_all('.kpi')}
    print('      cards:', cards)
    ok(len(cards) == 4, '4 cards')
    ok(pg.query_selector('#nav-visao.ativo') and pg.get_attribute('#nav-visao', 'data-tip') == 'Visão geral' and pg.get_attribute('#nav-matriz', 'data-tip') == 'Matriz', 'barra lateral com dica (Visão geral / Matriz)')
    pg.hover('#nav-matriz'); pg.wait_for_timeout(200)
    op = pg.evaluate("getComputedStyle(document.querySelector('#nav-matriz'), '::after').opacity")
    ok(op == '1', f'dica da aba aparece no hover (opacidade {op})')
    garagens = [t.inner_text() for t in pg.query_selector_all('#v-tabela tbody td.fx-g')]
    print('      garagens:', garagens)
    ok('Não informado' not in garagens and sum(1 for g in garagens if 'Sudeste' in g) == 1, f'{len(garagens)} Garagens/Empresas, sem "Não informado", Via Sudeste unificada')
    ok(num(total(pg, 'Veículos')) == nveic, f'Total de veículos = {nveic}')
    ok(num(total(pg, 'Veículos com 1+ câmera com falha')) == cards['veic'], 'Total "1+ câmera com falha" = card Veículos com falha')
    heads = [t.childNodes if False else t.inner_text() for t in pg.query_selector_all('#v-tabela thead th')]
    ok(all(f'Câm {c}' in ' '.join(heads) for c in range(21, 27)), 'colunas Câm 21 a Câm 26')
    ok(len(pg.query_selector_all('#v-graf svg path.ln')) == 3, 'gráfico Evolução diária com 3 linhas')
    # ordenação
    pg.click('#v-tabela th[data-o="vt"]'); pg.wait_for_timeout(150)
    vals = [num(t.inner_text()) for t in pg.query_selector_all('#v-tabela tbody tr td:nth-child(2)')]
    ok(vals == sorted(vals, reverse=True), 'ordenar por Veículos (decrescente)')
    pg.click('#v-tabela th[data-o="vt"]'); pg.wait_for_timeout(150)
    vals = [num(t.inner_text()) for t in pg.query_selector_all('#v-tabela tbody tr td:nth-child(2)')]
    ok(vals == sorted(vals), 'ordenar por Veículos (crescente)')
    pg.click('#v-tabela th[data-o="g"]'); pg.wait_for_timeout(150)
    for w, h in VIEWPORTS:
        pg.set_viewport_size({'width': w, 'height': h}); pg.wait_for_timeout(300)
        sem_rolagem(pg, f'Visão geral {w}x{h}')
        pg.screenshot(path=str(out / f'{pref}_vp_{w}x{h}_visao.png'))
    pg.set_viewport_size({'width': 1440, 'height': 900}); pg.wait_for_timeout(300)
    pg.screenshot(path=str(out / f'{pref}_1_visao_geral.png'))

    # manutenção
    pg.click('#v-manut'); pg.wait_for_timeout(250)
    m = {k: num(total(pg, k)) for k in ['Veículos atendidos', 'Reincidências', 'Atendimentos procedentes', 'Ocorrências solucionadas', 'Atendimentos improcedentes']}
    print('      manutenção (total):', m)
    ok(m['Veículos atendidos'] == 595 and m['Reincidências'] == 64 and m['Atendimentos procedentes'] == 282 and m['Ocorrências solucionadas'] == 71,
       'colunas de manutenção com os mesmos totais da análise (595 / 64 / 282 / 71)')
    ok(0 < m['Atendimentos improcedentes'] <= m['Veículos atendidos'] - m['Atendimentos procedentes'], 'Atendimentos improcedentes coerente')
    ok(all(pg.get_attribute(f'#v-tabela th[data-o="{k}"]', 'data-tip') for k in ['ma', 'mr', 'mp', 'ms', 'mi']), 'dicas com as definições nos cabeçalhos de manutenção')
    sem_rolagem(pg, 'Visão geral com manutenção')
    pg.screenshot(path=str(out / f'{pref}_2_visao_manutencao.png'))
    pg.click('#v-manut'); pg.wait_for_timeout(200)
    ok(pg.query_selector('#v-tabela th[data-o="ma"]') is None, 'desligar Manutenção remove as colunas')

    # Veículo | Câmera
    pg.click('#v-modo button[data-m="cam"]'); pg.wait_for_timeout(250)
    ok(num(total(pg, 'Câmeras funcionais')) == cards['on'] and num(total(pg, 'Câmeras com erro de SD card')) == cards['fa'] and num(total(pg, 'Câmeras 100% offline')) == cards['off'],
       'modo Câmera: totais = cards de câmeras')
    ok('Câmeras por dia' in pg.inner_text('#v-gsub'), 'gráfico segue o modo Câmera')
    pg.screenshot(path=str(out / f'{pref}_3_visao_camera.png'))
    pg.click('#v-modo button[data-m="veic"]'); pg.wait_for_timeout(250)

    # exportação Excel
    with pg.expect_download() as dl:
        pg.click('#v-xlsx')
    arq = out.parent / f'{pref}_garagens.xlsx'
    dl.value.save_as(str(arq))
    ws = openpyxl.load_workbook(arq).active
    linhas = [r for r in ws.iter_rows(values_only=True)]
    ok(any(r[0] == 'Total' and r[1] == nveic for r in linhas), f'Excel exportado ({dl.value.suggested_filename}) com linha Total')
    arq.unlink()

    # filtros e cards
    pg.select_option('#f-empresa', '*METROPOLE'); pg.select_option('#f-camera', '21'); pg.wait_for_timeout(500)
    esperado = sum(1 for v in frota['veiculos'] if v.get('e') is not None and 'METROPOLE' in meta['empresas'][v['e']].upper() and 21 in v['c'])
    ok(num(total(pg, 'Veículos')) == esperado and len(pg.query_selector_all('#v-tabela th.th-cam')) == 1, f'METROPOLE + câmera 21: {esperado} veículos, só a coluna Câm 21')
    pg.select_option('#f-empresa', ''); pg.select_option('#f-camera', ''); pg.wait_for_timeout(400)
    pg.type('#f-prefixo', '6800', delay=40); pg.wait_for_timeout(150)
    antes = num(total(pg, 'Veículos')); pg.wait_for_timeout(450)
    esp = sum(1 for v in frota['veiculos'] if '6800' in str(v['p']))
    ok(num(total(pg, 'Veículos')) == esp and antes != esp, f'prefixo "6800" com debounce: {esp}')
    pg.fill('#f-prefixo', ''); pg.wait_for_timeout(450)
    for k, esperado in (('on', 4961), ('fa', 1026), ('off', 1229), ('veic', cards['veic'])):
        pg.click(f'.kpi[data-card="{k}"]'); pg.wait_for_timeout(300)
        n = num(total(pg, 'Veículos'))
        bg = pg.eval_on_selector(f'.kpi[data-card="{k}"]', 'e => getComputedStyle(e).backgroundColor')
        ok(n == esperado and bg == 'rgb(230, 239, 234)', f'card {k} filtra tabela e gráfico ({n} veículos)')
        pg.click(f'.kpi[data-card="{k}"]'); pg.wait_for_timeout(300)
    ok(num(total(pg, 'Veículos')) == nveic, 'novo clique limpa o filtro do card')

    # detalhe da garagem
    linha = pg.query_selector('#v-tabela tbody tr[data-g="Via Sudeste"]')
    nv = num(linha.query_selector('td:nth-child(2)').inner_text())
    linha.click(); pg.wait_for_selector('#modal .gt-v')
    ok(len(pg.query_selector_all('#modal .gt-v tbody tr')) == nv, f'detalhe da garagem Via Sudeste: {nv} veículos')
    pg.click('#modal .gt-v th[data-o="p"]'); pg.wait_for_timeout(150)
    ps = [int(t.inner_text()) for t in pg.query_selector_all('#modal .gt-v tbody td.fx-g')]
    ok(ps == sorted(ps), 'lista de veículos ordenável')
    pg.click('#modal .gt-v th[data-o="sit"]'); pg.wait_for_timeout(150)
    alvo = pg.query_selector('#modal .gt-v tbody tr')
    p_alvo = alvo.get_attribute('data-p'); alvo.click(); pg.wait_for_selector('#modal .iv'); pg.wait_for_timeout(300)
    ok(f'Prefixo {p_alvo}' in pg.inner_text('#g-painel') and pg.query_selector('#g-painel .barra-tl'), f'veículo {p_alvo}: detalhe com linha do tempo')
    d0 = pg.inner_text('#g-painel .cab-p .sub'); pg.click('#pn-ant'); pg.wait_for_timeout(300)
    ok(pg.inner_text('#g-painel .cab-p .sub') != d0, 'setas trocam o dia no detalhe')
    sem_rolagem(pg, 'detalhe da garagem')
    pg.screenshot(path=str(out / f'{pref}_6_detalhe_garagem.png'))
    pg.keyboard.press('Escape'); pg.wait_for_timeout(200)
    ok(pg.query_selector('#modal') is None, 'Esc fecha o detalhe')

    # ---------------- Aba 2: Matriz ----------------
    pg.click('#nav-matriz'); pg.wait_for_selector('.st'); pg.wait_for_timeout(300)
    ok('#matriz' in pg.url and pg.query_selector('#nav-matriz.ativo'), 'aba Matriz')
    ok(pg.query_selector('#f-prefixo') is None and pg.query_selector('#f-de') is None, 'Matriz: filtros Empresa/Garagem, Câmera e Mês')
    meses = sorted({d[:7] for d in meta['dias']})
    ok(pg.input_value('#f-mes') == meses[-1] and len(pg.query_selector_all('#f-mes option')) == len(meses), f'Mês padrão = mais recente ({meses[-1]}), {len(meses)} meses')
    dias_mes = [d for d in meta['dias'] if d.startswith(meses[-1])]
    ok(len(pg.query_selector_all('.cel-h.dia-h')) == len(dias_mes), f'{len(dias_mes)} colunas de data no mês')
    leg = pg.inner_text('.legenda')
    ok(all(x in leg for x in ['Funcional', '100% Offline', 'Erro de SD card e/ou 1+ câmera com problema', 'Sem conexão', 'Manutenção']), 'legenda com os novos rótulos')
    alt = pg.evaluate("() => { const t = document.getElementById('m-tabela').getBoundingClientRect(); return [t.bottom, innerHeight]; }")
    ok(alt[1] - alt[0] < 90, f'matriz ocupa a altura disponível (fundo em {alt[0]:.0f} de {alt[1]})')
    total_m = num(pg.inner_text('#m-cont'))
    ok(total_m == nveic and len(pg.query_selector_all('#m-tabela .linha')) < 120, f'rolagem virtual: {len(pg.query_selector_all("#m-tabela .linha"))} linhas desenhadas de {total_m}')
    pg.eval_on_selector('#m-tabela', 'e => { e.scrollTop = e.scrollHeight; }'); pg.wait_for_timeout(300)
    vis = pg.evaluate("""() => { const t = document.getElementById('m-tabela').getBoundingClientRect(); const ls = document.querySelectorAll('#m-tabela .linha');
                          const c = ls[ls.length - 1].querySelector('.fx2').getBoundingClientRect(); return c.bottom <= t.bottom + 1 && c.top >= t.top; }""")
    cab_top = pg.evaluate("() => document.querySelector('#m-tabela .cel-h').getBoundingClientRect().top - document.getElementById('m-tabela').getBoundingClientRect().top")
    ok(vis and abs(cab_top) <= 2, 'fim da tabela alcançado e cabeçalho fixo')
    pg.eval_on_selector('#m-tabela', 'e => { e.scrollTop = 0; }'); pg.wait_for_timeout(200)
    pontos = pg.evaluate("""() => [...document.querySelectorAll('#m-tabela .linha')].map(l => [l.dataset.p,
                              [...l.querySelectorAll('.st')].map(c => [c.dataset.i, !!c.querySelector('.man')])])""")
    mv = {str(x['p']): x.get('mv', {}) for x in frota['veiculos']}
    ok(all(tem == bool(mv[pp].get(i)) for pp, cs in pontos for i, tem in cs), 'ponto azul = dias com manutenção')
    for w, h in VIEWPORTS:
        pg.set_viewport_size({'width': w, 'height': h}); pg.wait_for_timeout(300)
        sem_rolagem(pg, f'Matriz {w}x{h}')
        pg.screenshot(path=str(out / f'{pref}_vp_{w}x{h}_matriz.png'))
    pg.set_viewport_size({'width': 1440, 'height': 900}); pg.wait_for_timeout(300)
    pg.screenshot(path=str(out / f'{pref}_4_matriz.png'))
    alvo = pg.query_selector('#m-tabela .st.fa:has(.man), #m-tabela .st.off:has(.man)') or pg.query_selector('#m-tabela .st:has(.man)')
    p_alvo = alvo.evaluate("e => e.closest('.linha').dataset.p")
    alvo.click(); pg.wait_for_selector('#m-painel .iv'); pg.wait_for_timeout(350)
    txt = pg.inner_text('#m-painel')
    ok(f'Prefixo {p_alvo}' in txt and 'Manutenção' in txt and 'Ver registro completo' in txt, f'clique na célula: painel do prefixo {p_alvo} com manutenção')
    sem_rolagem(pg, 'Matriz com painel')
    pg.screenshot(path=str(out / f'{pref}_4b_matriz_painel.png'))
    pg.click('#pn-fechar'); pg.wait_for_timeout(200)

    # OS
    pg.click('#btn-os'); pg.wait_for_selector('#modal .os-form'); pg.wait_for_function("!document.querySelector('#os-res').textContent.includes('Calculando')")
    pg.select_option('#os-g', 'Via Sudeste'); pg.click('#os-p label:has(input[value="off"]) span'); pg.wait_for_timeout(800)
    ok(pg.is_checked('#os-p input[value="off"]') and pg.is_checked('#os-p input[value="prob"]'), 'OS: seleção múltipla de problemas')
    n_os = num(pg.inner_text('#os-res b'))
    pg.screenshot(path=str(out / f'{pref}_5_modal_os.png'))
    with pg.expect_download() as dl:
        pg.click('#os-gerar')
    arq_os = out.parent / 'OS_exemplo.xlsx'
    dl.value.save_as(str(arq_os))
    wb = openpyxl.load_workbook(arq_os)
    ws = wb['Ordem de Serviço']
    cab = [c.value for c in ws[7]]
    print('      OS:', dl.value.suggested_filename, wb.sheetnames, cab)
    dados = [r for r in ws.iter_rows(min_row=8, values_only=True) if r[0] is not None]
    ok(wb.sheetnames == ['Ordem de Serviço', 'Legenda e instruções'] and cab[:3] == ['Prioridade', 'Garagem', 'Prefixo'] and cab[-4:] == ['Ação realizada', 'Data do atendimento', 'Técnico', 'Assinatura'],
       'OS: abas e colunas')
    ok(len(dados) == n_os and [r[0] for r in dados] == list(range(1, n_os + 1)), f'OS: {n_os} veículos numerados por prioridade')
    di = cab.index('Dias com problema')
    ok(all(dados[k][di] >= dados[k + 1][di] for k in range(len(dados) - 1)), 'OS: ordenada pelos dias com problema')
    ok(ws.freeze_panes == 'D8' and ws.auto_filter.ref and ws.page_setup.orientation == 'landscape' and ws.page_setup.paperSize in (9, '9') and ws.sheet_properties.pageSetUpPr.fitToPage,
       'OS: cabeçalho congelado, filtros, A4 paisagem ajustada à largura')
    pg.keyboard.press('Escape')
    # celular: painel como gaveta
    pg.set_viewport_size({'width': 390, 'height': 844}); pg.wait_for_timeout(300)
    pg.query_selector('#m-tabela .st:not(.nd)').click(); pg.wait_for_selector('#m-painel .iv'); pg.wait_for_timeout(300)
    ok(pg.eval_on_selector('#m-painel .painel', 'e => getComputedStyle(e).position') == 'fixed', '< 1024 px: painel como gaveta')
    sem_rolagem(pg, 'celular com painel')
    b.close()
print('ERROS DE CONSOLE:', erros or 'nenhum')
print('FALHAS:', falhas or 'nenhuma')
sys.exit(1 if falhas or erros else 0)
