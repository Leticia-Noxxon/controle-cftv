"""Relatório CFTV diário (planilha 'Relatório CFTV - DD.MM.AAAA.xlsx'): retrato consolidado do dia, usado como
fonte complementar da situação atual (28/09/2026). Mesmas regras do projeto anterior:
'-' = câmera não instalada (ignorada); vazio = coluna inexistente; OFFLINE; ONLINE com SD/Login/Gravação."""
import datetime as dt
import re

import openpyxl

from . import config

_RX = re.compile(r'SD:\s*(\w+).*?Login:\s*(\w+).*?Grava\w*:\s*(\w+)', re.IGNORECASE)


def interpretar(txt):
    if txt is None or str(txt).strip() == '':
        return '.'
    t = str(txt).strip()
    if t == '-':
        return '-'
    if t.upper().startswith('OFFLINE'):
        return 'O'
    m = _RX.search(t)
    if t.upper().startswith('ONLINE') and m:
        b = (1 if m.group(1).lower() == 'error' else 0) + (2 if m.group(2).lower() == 'error' else 0) + (4 if m.group(3).lower() == 'error' else 0)
        return 'N' if b == 0 else str(b)
    return '?'


def ler():
    out = []
    for path in config.arquivos_relatorio():
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        aba = None
        for ws in wb.worksheets:
            cab = [str(c).strip() if c is not None else '' for c in next(ws.iter_rows(max_row=1, values_only=True))]
            if 'Prefixo' in cab and any(c.startswith('Câmera 2') for c in cab):
                aba = ws
                break
        if aba is None:
            continue
        ic = {n: cab.index(f'Câmera {n}') for n in config.CAMERAS if f'Câmera {n}' in cab}
        regs = {}
        data = None
        for n, row in enumerate(aba.iter_rows(min_row=2, values_only=True), start=2):
            try:
                p = int(row[cab.index('Prefixo')])
            except (TypeError, ValueError):
                continue
            d = row[cab.index('Data')]
            data = d.date() if isinstance(d, dt.datetime) else data
            regs[p] = {'linha': n, 'empresa': str(row[cab.index('Empresa')] or '').strip(),
                       'status': row[cab.index('Status')] if 'Status' in cab else None,
                       'cams': {c: interpretar(row[i]) for c, i in ic.items()},
                       'textos': {c: row[i] for c, i in ic.items()}}
        out.append({'arquivo': path.name, 'aba': aba.title, 'abas': wb.sheetnames, 'data': data.isoformat() if data else None, 'registros': regs})
        wb.close()
    return out
