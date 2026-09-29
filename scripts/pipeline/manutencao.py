"""Leitura do formulário de manutenção (Revisão_CFTV*.xlsx).

- Lê a linha inteira de cada resposta (todas as colunas).
- O arquivo mistura dois layouts: as respostas mais novas têm a coluna 'Data de Envio' (e 'ID do Envio'/'IP do Envio');
  um bloco de respostas antigas foi colado com as colunas deslocadas uma posição à esquerda (layout antigo, sem
  'Data de Envio'). Essas linhas são realinhadas pelo layout antigo; quando são cópia idêntica de uma resposta já
  existente, são marcadas como duplicadas e não entram na análise (continuam na auditoria).
- Cada célula de problema/ação pode conter vários itens separados por quebra de linha: todos são separados e
  normalizados para um vocabulário padrão (sem perder o texto original).
"""
import datetime as dt
import re
import unicodedata
from collections import Counter

import openpyxl

from . import config

MESES = {'janeiro': 1, 'fevereiro': 2, 'março': 3, 'marco': 3, 'abril': 4, 'maio': 5, 'junho': 6, 'julho': 7,
         'agosto': 8, 'setembro': 9, 'outubro': 10, 'novembro': 11, 'dezembro': 12}
_RX_DATA = re.compile(r'^\s*[^,]+,\s*([A-Za-zçÇ]+)\s+(\d{1,2}),\s*(\d{4})\s+(\d{1,2}):(\d{2})')
_RX_COL = re.compile(r'^(Problemas Detectados|Ações)\s*:\s*C[âa]m\.?\s*(.+)$', re.IGNORECASE)

LAYOUT_ANTIGO = ['Prefixo', 'Nome', 'Data', 'ID', 'Garagem', 'Tecnologia', 'Câmera Utilizada',
                 'Problemas Detectados: Câm. Frontal', 'Ações: Câm. Frontal', 'Problemas Detectados: Câm. Frente',
                 'Ações: Câm. Frente', 'Problemas Detectados: Câm. Corredor 1', 'Problemas Detectados: Câm. Corredor 2',
                 'Ações: Câm. Corredor 2', 'Switch Utilizado', 'Problemas Detectados: Câm. Corredor 3',
                 'Cartão de Memória Utilizado', 'Ações: Câm. Corredor 1', 'Ações: Câm. Corredor 3',
                 'Problemas Detectados: Câm. Corredor 4', 'Ações: Câm. Corredor 4', 'Observações']
COLUNAS_NAO_PUBLICADAS = ['IP do Envio']   # endereço IP de quem enviou: não é necessário para a análise

# Vocabulário padrão (opções do formulário) -> categoria
PROBLEMAS = {
    'Sem gravação de imagens': 'Gravação', 'Configuração incorreta da câmera': 'Configuração',
    'Câmera mal reposicionada': 'Posicionamento', 'Câmera não fixada corretamente': 'Posicionamento',
    'Câmera inoperante': 'Câmera', 'Câmera travada': 'Câmera', 'Câmera com infiltração': 'Câmera',
    'Câmera desligada manualmente': 'Câmera', 'Câmera removida por terceiros': 'Vandalismo/terceiros',
    'Câmera com sinal de vandalismo': 'Vandalismo/terceiros', 'Câmera sem cartão de memória': 'Cartão SD',
    'Cabeamento rompido/danificado': 'Cabeamento', 'Falha na conexão dos cabos': 'Cabeamento',
    'RJ45 crimpado incorretamente': 'Cabeamento', 'Switch apresentando falha': 'Switch',
    'Conexões do switch incorretas': 'Switch', 'Switch não fixado corretamente': 'Switch',
    'Veículo sem bateria': 'Energia', 'Fusível queimado': 'Energia',
}
ACOES = {
    'Normalização da gravação de imagens': 'Gravação', 'Correção e ajuste da configuração da câmera': 'Configuração',
    'Reposicionamento adequado da câmera': 'Posicionamento', 'Fixação adequada da câmera': 'Posicionamento',
    'Reinicialização da câmera': 'Câmera', 'Ativação da câmera': 'Câmera', 'Substituição da câmera defeituosa': 'Câmera',
    'Instalação da câmera removida': 'Câmera', 'Câmera encaminhada para manutenção': 'Câmera',
    'Inserção ou substituição do cartão de memória': 'Cartão SD', 'Substituição do cabeamento': 'Cabeamento',
    'Reparo do cabeamento': 'Cabeamento', 'Crimpagem correta do conector RJ45': 'Cabeamento',
    'Correção na conexão dos cabos': 'Cabeamento', 'Substituição do switch': 'Switch',
    'Correção das conexões do switch': 'Switch', 'Fixação adequada do switch': 'Switch', 'Correção da falha no switch': 'Switch',
    'Substituição do fusível': 'Energia',
}
SEM_PROBLEMA = 'Nenhuma anomalia identificada'
SEM_ACAO = 'Nenhuma ação realizada'

# Itens encontrados no texto livre (Observações): termo -> (tipo, item padronizado)
TEXTO_LIVRE = [
    (r'formata\w*', ('acao', 'Formatação do cartão SD')),
    (r'(troca|substitui\w*)\s+(d[eo]s?\s+)?(cart[ãa]o|sd)|pen\s*drive\s+pelo\s+cart', ('acao', 'Troca do cartão SD')),
    (r'configura\w*', ('acao', 'Configuração da câmera')),
    (r'ajuste\s+de\s+imagem|reposicion\w*|posi[çc][ãa]o|posicionad\w*', ('acao', 'Ajuste de imagem/posição')),
    (r'(troca|substitu\w*)\s+(d[ae]\s+)?c[âa]m|c[âa]mera\s+substitu\w*|substitu[ií]da\s+c[âa]mera', ('acao', 'Substituição da câmera')),
    (r'cabo|cabeamento|rj ?45|crimp\w*', ('acao', 'Cabeamento')),
    (r'switch|\bpoe\b', ('acao', 'Switch/PoE')),
    (r'\bucp\b|\btdm\b', ('acao', 'UCP/TDM')),
    (r'em curto|curto[ -]circuito', ('problema', 'Câmera em curto')),
    (r'inoperante|inativa|n[ãa]o liga|com defeito|defeituos\w*', ('problema', 'Câmera inoperante/defeito')),
    (r'infiltra\w*', ('problema', 'Infiltração')),
    (r'vandal\w*|retirad\w* por terceiros|removid\w* por terceiros', ('problema', 'Vandalismo/terceiros')),
    (r'100\s*%\s*online', ('resultado', '100% ONLINE (informado pelo técnico)')),
]
PENDENCIA = [r'\bpendente\b', r'\bpend[êe]ncias?\b', r'\baguardand\w*', r'n[ãa]o foi poss[íi]vel', r'\bfalta\b',
             r'\bfaltando\b', r'fazer a troca', r'solicitad\w*', r'\bnecess[áa]ri[oa]\b']
_RX_CAM = re.compile(r'\bc[aâ]m(?:[eêa]ras?|aras?)?\b\.?[\s,:]*((?:2[1-6])(?:\s*(?:,|;|e|&|/|\s)\s*2[1-6])*)\b', re.IGNORECASE)
_RX_C = re.compile(r'\bC([1-6])\b')
_RX_HEX12 = re.compile(r'\b[0-9A-F]{12}\b')
_RX_SERIE = re.compile(r'\b(?=[0-9A-Z]*\d)(?=[0-9A-Z]*[A-Z])[0-9A-Z]{10,}\b')


def _sa(s):
    return ''.join(c for c in unicodedata.normalize('NFD', str(s)) if unicodedata.category(c) != 'Mn')


def _chave(s):
    return re.sub(r'\s+', ' ', _sa(s).strip().lower().rstrip('.;'))


_PROB_K = {_chave(k): k for k in PROBLEMAS}
_ACAO_K = {_chave(k): k for k in ACOES}


def vazio(v):
    return v is None or (isinstance(v, str) and not v.strip())


def parse_data(txt):
    m = _RX_DATA.match(str(txt or ''))
    if not m or m.group(1).lower() not in MESES:
        return None
    return dt.datetime(int(m.group(3)), MESES[m.group(1).lower()], int(m.group(2)), int(m.group(4)), int(m.group(5)))


def classificar_item(txt):
    """Retorna (tipo, item_padronizado, categoria). tipo: problema | acao | nenhum_problema | nenhuma_acao | quantidade | texto"""
    k = _chave(txt)
    if k == _chave(SEM_PROBLEMA):
        return 'nenhum_problema', SEM_PROBLEMA, None
    if k == _chave(SEM_ACAO):
        return 'nenhuma_acao', SEM_ACAO, None
    if k in _PROB_K:
        p = _PROB_K[k]
        return 'problema', p, PROBLEMAS[p]
    if k in _ACAO_K:
        a = _ACAO_K[k]
        return 'acao', a, ACOES[a]
    if re.fullmatch(r'\d+', k):
        return 'quantidade', k, None
    return 'texto', str(txt).strip(), None


def itens(v):
    return [] if vazio(v) else [x.strip() for x in str(v).split('\n') if x.strip()]


def cameras_no_texto(txt):
    if vazio(txt):
        return []
    out = []
    for g in _RX_CAM.findall(str(txt)):
        for n in re.findall(r'2[1-6]', g):
            if int(n) not in out:
                out.append(int(n))
    return sorted(out)


def analisar_texto(txt):
    """Estrutura o texto livre: trechos por câmera citada, itens reconhecidos, pendências, série/MAC."""
    res = {'por_camera': {}, 'itens': [], 'notacao_c': [], 'pendencias': [], 'equipamento': equipamentos(txt)}
    if vazio(txt):
        return res
    texto = str(txt)
    res['notacao_c'] = sorted(set('C' + n for n in _RX_C.findall(texto)))
    vistos = set()
    for linha in [l.strip() for l in re.split(r'\n|(?<=[.;])\s+', texto) if l.strip()]:
        low = _sa(linha).lower()
        cams = cameras_no_texto(linha)
        achados = []
        for rx, (tipo, item) in TEXTO_LIVRE:
            if re.search(rx, low, re.IGNORECASE):
                achados.append({'tipo': tipo, 'item': item})
                if (tipo, item) not in vistos:
                    vistos.add((tipo, item))
                    res['itens'].append({'tipo': tipo, 'item': item, 'trecho': linha})
        for c in cams:
            d = res['por_camera'].setdefault(str(c), {'trechos': [], 'itens': []})
            d['trechos'].append(linha)
            for a in achados:
                if a not in d['itens']:
                    d['itens'].append(a)
        for p in PENDENCIA:
            m = re.search(p, low, re.IGNORECASE)
            if m:
                res['pendencias'].append({'trecho': linha, 'termo': m.group(0)})
                break
    return res


def equipamentos(txt):
    res = {'instalado': {'series': [], 'macs': []}, 'retirado': {'series': [], 'macs': []}, 'nao_especificado': {'series': [], 'macs': []}}
    if vazio(txt):
        return res
    estado = 'nao_especificado'
    for linha in str(txt).split('\n'):
        low = _sa(linha).lower()
        if 'retirad' in low:
            estado = 'retirado'
        elif 'instalad' in low:
            estado = 'instalado'
        up = linha.upper()
        if 'MAC' in up:
            for mac in _RX_HEX12.findall(up):
                if mac not in res[estado]['macs']:
                    res[estado]['macs'].append(mac)
            continue
        for t in _RX_SERIE.findall(up):
            if t not in res[estado]['series'] and not re.fullmatch(r'\d+', t):
                res[estado]['series'].append(t)
    return res


def ler(path=None):
    path = path or config.arquivo_manutencao()
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    abas = wb.sheetnames
    ws = wb[abas[0]]
    linhas = list(ws.iter_rows(values_only=True))
    wb.close()
    cab = [str(c).strip() if c is not None else '' for c in linhas[0]]
    idx = {c: i for i, c in enumerate(cab)}
    respostas, auditoria = [], []
    for n, row in enumerate(linhas[1:], start=2):
        if all(vazio(v) for v in row):
            continue
        d = {c: row[i] for c, i in idx.items() if i < len(row)}
        layout = 'atual'
        # linha no layout antigo, deslocada 1 coluna: 'ID do Envio' vazio, 'Data' numérico e 'Data de Envio' com data por extenso
        if vazio(d.get('ID do Envio')) and isinstance(d.get('Data'), (int, float)) and parse_data(d.get('Data de Envio')):
            d = {c: (row[i] if i < len(row) else None) for i, c in enumerate(LAYOUT_ANTIGO)}
            layout = 'antigo_realinhado'
        respostas.append({'linha_excel': n, 'layout': layout, 'campos': d})
    # duplicatas: linha realinhada idêntica a uma resposta no layout atual (mesmo prefixo, data, ID e mesmas respostas)
    chave = lambda c: (c.get('Prefixo'), str(c.get('Data')), c.get('ID'))
    atuais = {chave(r['campos']): r for r in respostas if r['layout'] == 'atual'}
    comparar = [c for c in LAYOUT_ANTIGO]
    for r in respostas:
        r['duplicada_de'] = None
        if r['layout'] != 'atual':
            a = atuais.get(chave(r['campos']))
            if a and all((str(a['campos'].get(c) or '').strip() == str(r['campos'].get(c) or '').strip()) for c in comparar):
                r['duplicada_de'] = a['linha_excel']
    info = {'arquivo': path.name, 'abas': abas, 'aba_utilizada': abas[0], 'colunas': cab, 'linhas': len(respostas),
            'realinhadas': sum(r['layout'] != 'atual' for r in respostas),
            'duplicadas': sum(r['duplicada_de'] is not None for r in respostas)}
    return respostas, info


def processar(respostas, prefixos_monitorados):
    forms = []
    nomes = Counter(str(r['campos'].get('Nome') or '').strip() for r in respostas if r['duplicada_de'] is None)
    por_lower = {}
    for nome, q in nomes.items():
        por_lower.setdefault(nome.lower(), []).append((q, nome))
    # mesma pessoa com grafias que diferem só em maiúsculas/minúsculas: usa a grafia com iniciais maiúsculas (se houver),
    # senão a mais frequente (regra herdada do projeto anterior: 'Abner melo' -> 'Abner Melo')
    def _melhor(v):
        titulo = [n for _, n in v if n == n.title()]
        return titulo[0] if titulo else max(v)[1]
    padrao_nome = {n: _melhor(v) for v in por_lower.values() for _, n in v}
    for r in respostas:
        c = r['campos']
        alertas = []
        data = parse_data(c.get('Data'))
        if data is None:
            alertas.append('Data do formulário não reconhecida')
        prefixo = c.get('Prefixo')
        try:
            prefixo = int(prefixo)
        except (TypeError, ValueError):
            alertas.append('Prefixo inválido')
            prefixo = None
        if prefixo is not None and prefixo not in prefixos_monitorados:
            alertas.append('Prefixo não aparece no monitoramento')
        if r['layout'] != 'atual':
            alertas.append('Linha no layout antigo (colunas deslocadas) — realinhada')
        posicoes = []
        for col in c:
            m = _RX_COL.match(col or '')
            if not m:
                continue
            tipo_col = 'problemas' if m.group(1).lower().startswith('problema') else 'acoes'
            pos = re.sub(r'\s+', ' ', m.group(2).strip().upper())
            pos = 'CORREDOR 1' if pos == 'CORREDOR' else pos
            cam = config.POSICAO_PARA_CAMERA.get(pos)
            p = next((x for x in posicoes if x['posicao'] == pos), None)
            if p is None:
                p = {'posicao': pos, 'camera': cam, 'problemas': [], 'acoes': [], 'outros': [], 'original': {}}
                posicoes.append(p)
            p['original'][col] = c[col]
            for it in itens(c[col]):
                t, item, cat = classificar_item(it)
                if t == 'problema':
                    p['problemas'].append({'item': item, 'categoria': cat, 'coluna_trocada': tipo_col != 'problemas'})
                elif t == 'acao':
                    p['acoes'].append({'item': item, 'categoria': cat, 'coluna_trocada': tipo_col != 'acoes'})
                elif t in ('nenhum_problema', 'nenhuma_acao'):
                    pass
                else:
                    p['outros'].append({'texto': item, 'coluna': col})
        for p in posicoes:
            if any(x['coluna_trocada'] for x in p['problemas'] + p['acoes']):
                alertas.append(f"{p['posicao']}: item de problema/ação registrado na coluna trocada")
        obs = c.get('Observações')
        texto = analisar_texto(obs)
        cams_form = sorted({p['camera'] for p in posicoes if p['camera'] and (p['problemas'] or p['acoes'])})
        nome = str(c.get('Nome') or '').strip()
        f = {
            'id': r['linha_excel'], 'linha_excel': r['linha_excel'], 'layout': r['layout'], 'duplicada_de': r['duplicada_de'],
            'prefixo': prefixo, 'tecnico': padrao_nome.get(nome, nome), 'tecnico_original': nome,
            'data_texto': c.get('Data'), 'datahora': data.isoformat(timespec='minutes') if data else None,
            'data': data.date().isoformat() if data else None, 'hora': data.strftime('%H:%M') if data else None,
            'data_envio': c.get('Data de Envio') if r['layout'] == 'atual' else None,
            'id_formulario': c.get('ID'), 'garagem': c.get('Garagem'), 'tecnologia': c.get('Tecnologia'),
            'quantidades': {k: (None if vazio(c.get(k)) else str(c.get(k))) for k in ['Câmera Utilizada', 'Switch Utilizado', 'Cartão de Memória Utilizado']},
            'posicoes': posicoes, 'cameras_formulario': cams_form,
            'cameras_texto': cameras_no_texto(obs), 'texto': texto, 'observacoes': None if vazio(obs) else str(obs),
            'id_envio': c.get('ID do Envio'), 'alertas': alertas,
            'respostas': [{'coluna': k, 'valor': v} for k, v in c.items() if not vazio(v) and k not in COLUNAS_NAO_PUBLICADAS],
        }
        f['problemas'] = sorted({x['item'] for p in posicoes for x in p['problemas']})
        f['acoes'] = sorted({x['item'] for p in posicoes for x in p['acoes']})
        forms.append(f)
    return forms
